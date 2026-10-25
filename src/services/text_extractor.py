from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class TextExtractionError(ValueError):
    """Raised when a document's text cannot be extracted."""


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from PDF, DOCX or UTF-8 TXT document bytes."""
    extension = Path(filename or "").suffix.lower()

    if extension not in {".pdf", ".docx", ".txt"}:
        raise TextExtractionError(
            "Only PDF, DOCX and TXT files are supported."
        )

    if not content:
        raise TextExtractionError("The document is empty.")

    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")

        elif extension == ".docx":
            document = Document(BytesIO(content))
            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Include text contained in tables too.
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            parts.append(cell.text)

            text = "\n".join(parts)

        else:  # PDF
            reader = PdfReader(BytesIO(content), strict=False)

            if reader.is_encrypted:
                raise TextExtractionError(
                    "The PDF is password-protected. "
                    "Provide an accessible, unprotected PDF."
                )

            text = "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

    except TextExtractionError:
        raise
    except Exception as exc:
        raise TextExtractionError(
            "The document could not be read. "
            "It may be corrupted or password-protected."
        ) from exc

    text = text.replace("\x00", "").strip()

    if not text:
        raise TextExtractionError(
            "No extractable text was found. "
            "Scanned PDFs are unsupported because OCR is out of scope."
        )

    return text